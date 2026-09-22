from __future__ import annotations
import csv
import io
from dataclasses import dataclass
from typing import BinaryIO

REQUIRED_COLUMNS = ("id", "inttype", "name")
ALL_COLUMNS = ("id", "intparentid", "inttype", "name", "contact", "phone", "phoneext", "address1", "address2", "city", "state", "zip", "reference")
MAX_LENGTHS = {"name": 50, "contact": 100, "phone": 50, "phoneext": 50, "address1": 255, "address2": 255, "city": 100, "state": 5, "zip": 10, "reference": 20}

@dataclass
class ImportResult:
    imported: int
    updated: int
    skipped: int
    errors: list[dict]
    columns: list[str]


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _validate(row: dict[str, str], row_number: int) -> list[dict]:
    errors = []
    for column in REQUIRED_COLUMNS:
        if not row.get(column): errors.append({"row": row_number, "column": column, "message": "Required value is missing"})
    for column, maximum in MAX_LENGTHS.items():
        if len(row.get(column, "")) > maximum: errors.append({"row": row_number, "column": column, "message": f"Maximum length is {maximum}"})
    for column in ("id", "intparentid", "inttype"):
        value = row.get(column, "")
        if value and not value.isdigit(): errors.append({"row": row_number, "column": column, "message": "Must be an integer"})
    return errors


def read_rows(filename: str, content: bytes) -> tuple[list[dict[str, str]], list[str]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        headers = [(_clean(header)).lower() for header in (reader.fieldnames or [])]
        return [{header: _clean(row.get(original)) for original, header in zip(reader.fieldnames or [], headers)} for row in reader], headers
    if suffix in {"xls", "xlsx"}:
        if suffix == "xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content)
            sheet = workbook.sheet_by_index(0)
            values = [[_clean(value) for value in sheet.row_values(index)] for index in range(sheet.nrows)]
        else:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = [[_clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        headers = [value.lower() for value in (values[0] if values else [])]
        return [dict(zip(headers, row)) for row in values[1:]], headers
    raise ValueError("Only .csv, .xls, and .xlsx files are supported")
